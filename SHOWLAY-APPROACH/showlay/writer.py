"""Excel writers for template-shaped extraction output and review sidecars."""
from __future__ import annotations

import shutil
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .schema import FIELD_HEADER_ALIASES, Row, map_template_columns

_GREEN = PatternFill(start_color="D6F5D6", end_color="D6F5D6", fill_type="solid")
_YELLOW = PatternFill(start_color="FFF2C7", end_color="FFF2C7", fill_type="solid")
_RED = PatternFill(start_color="FBD3D3", end_color="FBD3D3", fill_type="solid")
_QTYPE_NAMES = {"questiontype", "question type"}
_ASSESSMENT_SHEET = "Assessment"
_UNSECTIONED_CONTENT = "unsectioned content"

_TEMPLATE_OUTPUT_FIELDS = [
    field_name
    for field_name in FIELD_HEADER_ALIASES
    if field_name
    not in {
        "external_id",
        "branching_source",
        "source_page",
        "risk_level",
        "review_notes",
    }
]


def _confidence_fill(c: float) -> PatternFill | None:
    if c >= 0.9:
        return _GREEN
    if c >= 0.7:
        return _YELLOW
    return _RED


def _risk_level(row: Row) -> str:
    if row.confidence < 0.65:
        return "high"
    if row.confidence < 0.85:
        return "medium"
    return "low"


def _safe_excel_value(value):
    if isinstance(value, str):
        return ILLEGAL_CHARACTERS_RE.sub("", value)
    return value


def _excel_value(field_name: str, row: Row):
    value = getattr(row, field_name, "")
    if field_name == "section" and isinstance(value, str):
        if " ".join(value.split()).strip().lower() == _UNSECTIONED_CONTENT:
            return ""
    return _safe_excel_value(value)


def _find_header_row(ws) -> int:
    """Find the row containing the template's QuestionType / Question Type header."""
    for row_idx in range(1, min(ws.max_row, 40) + 1):
        for col_idx in range(1, min(ws.max_column, 50) + 1):
            value = ws.cell(row=row_idx, column=col_idx).value
            if isinstance(value, str) and " ".join(value.split()).strip().lower() in _QTYPE_NAMES:
                return row_idx
    raise ValueError(f"Could not find QuestionType header row in sheet {ws.title!r}.")


def _clear_data_rows(ws, header_row: int, max_col: int) -> None:
    for row_idx in range(header_row + 1, ws.max_row + 1):
        for col_idx in range(1, max_col + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.value = None
            cell.fill = PatternFill(fill_type=None)


def _is_metadata_label(value) -> bool:
    return isinstance(value, str) and value.strip().endswith(":")


def _clear_metadata_values(ws, header_row: int) -> None:
    for row_idx in range(1, header_row):
        values = [ws.cell(row=row_idx, column=col_idx).value for col_idx in range(1, ws.max_column + 1)]
        if not any(_is_metadata_label(value) for value in values):
            continue
        for col_idx, value in enumerate(values, start=1):
            if value in (None, "") or _is_metadata_label(value):
                continue
            ws.cell(row=row_idx, column=col_idx).value = None


def write_workbook(template_path: str, out_path: str, rows: list[Row],
                   sheet_name: str | None = None) -> str:
    """Clone the CHOICES HIP workbook and write rows into the old Assessment sheet."""
    Path(out_path).parent.mkdir(parents=True, exist_ok=True)
    shutil.copyfile(template_path, out_path)
    wb = load_workbook(out_path)
    target_sheet = sheet_name or _ASSESSMENT_SHEET
    if target_sheet not in wb.sheetnames:
        raise ValueError(f"Template workbook must contain an {target_sheet!r} sheet.")
    ws = wb[target_sheet]
    wb.active = wb.sheetnames.index(target_sheet)
    if target_sheet == _ASSESSMENT_SHEET and "Assessment v2" in wb.sheetnames:
        del wb["Assessment v2"]
        wb.active = wb.sheetnames.index(target_sheet)

    header_row = _find_header_row(ws)
    header_values = {col_idx: ws.cell(row=header_row, column=col_idx).value
                     for col_idx in range(1, ws.max_column + 1)}
    field_to_col = map_template_columns(header_values)
    max_col = max(list(field_to_col.values()) + [ws.max_column])

    _clear_metadata_values(ws, header_row)
    _clear_data_rows(ws, header_row, max_col)

    for row_offset, row in enumerate(rows, start=1):
        sheet_row = header_row + row_offset
        for field_name in _TEMPLATE_OUTPUT_FIELDS:
            col_idx = field_to_col.get(field_name)
            if col_idx is None:
                continue
            value = _excel_value(field_name, row)
            if field_name == "sequence":
                value = value if value is not None else ""
            cell = ws.cell(row=sheet_row, column=col_idx, value=value)
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    wb.save(out_path)
    return out_path


def write_review_sidecar(out_path: str, rows: list[Row]) -> str:
    """Compact human-review queue: rows sorted by confidence asc, with reasons + source page."""
    Path(out_path).parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "Review Queue"
    headers = ["Confidence", "Risk", "Page", "Sequence", "QuestionType", "Question Text",
               "Section", "Answer Text", "Answer Validation", "Branching Logic", "Question Rule", "Review Reasons"]
    for c, h in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.font = Font(bold=True)

    sorted_rows = sorted(rows, key=lambda r: (r.confidence, r.sequence or 0))
    for i, r in enumerate(sorted_rows, start=2):
        ws.cell(row=i, column=1, value=r.confidence)
        ws.cell(row=i, column=2, value=_safe_excel_value(_risk_level(r)))
        ws.cell(row=i, column=3, value=r.page or "")
        ws.cell(row=i, column=4, value=r.sequence or "")
        ws.cell(row=i, column=5, value=_safe_excel_value(r.question_type))
        ws.cell(row=i, column=6, value=_safe_excel_value(r.question_text))
        ws.cell(row=i, column=7, value=_excel_value("section", r))
        ws.cell(row=i, column=8, value=_safe_excel_value(r.answer_text))
        ws.cell(row=i, column=9, value=_safe_excel_value(r.answer_validation))
        ws.cell(row=i, column=10, value=_safe_excel_value(r.branching_logic))
        ws.cell(row=i, column=11, value=_safe_excel_value(r.question_rule))
        ws.cell(row=i, column=12, value=_safe_excel_value("; ".join(r.review_reasons)))
        fill = _confidence_fill(r.confidence)
        if fill:
            for c in range(1, len(headers) + 1):
                ws.cell(row=i, column=c).fill = fill

    for c, w in enumerate([10, 10, 8, 9, 14, 60, 24, 40, 28, 42, 34, 44], start=1):
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.row_dimensions[1].height = 22
    wb.save(out_path)
    return out_path
