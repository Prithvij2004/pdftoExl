from __future__ import annotations

import sys
from pathlib import Path

from openpyxl import load_workbook

SHOWLAY_DIR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(SHOWLAY_DIR))

from showlay.schema import Row  # noqa: E402
from showlay.writer import write_review_sidecar, write_workbook  # noqa: E402


def test_write_workbook_uses_choices_assessment_template_sheet(tmp_path):
    row = Row(
        sequence=1,
        section="Section A",
        question_rule="Display if Q2 = Yes",
        question_type="Text Box",
        question_text="Applicant Name",
        answer_validation="default characters = 100",
        page=1,
        confidence=0.91,
    )
    row.external_id = "A0100"

    template_path = SHOWLAY_DIR / "support_docs" / "choices-safety-determination-hip-workbook.xlsx"
    out_path = tmp_path / "out.xlsx"
    write_workbook(str(template_path), str(out_path), [row])

    wb = load_workbook(out_path, data_only=True)
    ws = wb.active
    headers = [ws.cell(row=14, column=idx).value for idx in range(1, 29)]

    assert ws.title == "Assessment"
    assert "Assessment v2" not in wb.sheetnames
    assert len(headers) == 28
    assert headers[10:16] == [
        "Sequence",
        "Question Rule",
        "QuestionType",
        "Question Text",
        "Branching Logic",
        "Answer Text",
    ]
    assert "External ID" not in headers
    assert ws["A15"].value == "Section A"
    assert ws["K15"].value == 1
    assert ws["L15"].value == "Display if Q2 = Yes"
    assert ws["M15"].value == "Text Box"
    assert ws["N15"].value == "Applicant Name"


def test_write_workbook_blanks_unsectioned_content_section(tmp_path):
    row = Row(
        sequence=1,
        section="Unsectioned Content",
        question_type="Text Box",
        question_text="Applicant Name",
        answer_validation="default characters = 100",
        confidence=0.91,
    )

    template_path = SHOWLAY_DIR / "support_docs" / "choices-safety-determination-hip-workbook.xlsx"
    out_path = tmp_path / "out.xlsx"
    write_workbook(str(template_path), str(out_path), [row])

    wb = load_workbook(out_path, data_only=True)
    ws = wb.active

    assert ws["A15"].value is None
    assert ws["N15"].value == "Applicant Name"


def test_writers_strip_illegal_excel_characters(tmp_path):
    row = Row(
        sequence=1,
        section="Section A",
        question_type="Display",
        question_text="Bad\x00question",
        branching_logic="Skip\x00logic",
        confidence=0.91,
        review_reasons=["Bad\x00reason"],
    )
    template_path = SHOWLAY_DIR / "support_docs" / "choices-safety-determination-hip-workbook.xlsx"
    out_path = tmp_path / "out.xlsx"
    review_path = tmp_path / "review.xlsx"

    write_workbook(str(template_path), str(out_path), [row])
    write_review_sidecar(str(review_path), [row])

    wb = load_workbook(out_path, data_only=True)
    assert wb.active["N15"].value == "Badquestion"
    assert wb.active["O15"].value == "Skiplogic"

    review_wb = load_workbook(review_path, data_only=True)
    assert review_wb.active.cell(row=2, column=6).value == "Badquestion"
    assert review_wb.active.cell(row=2, column=12).value == "Badreason"
