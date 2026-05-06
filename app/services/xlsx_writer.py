from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.services.merge_resolve import WORKBOOK_COLUMNS, FinalRow


XLSX_WRITER_VERSION = "xlsx-writer-v1"

SHEET_NAME = "Form"

HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL_COLOR = "1F3864"
HEADER_FILL = PatternFill("solid", fgColor=HEADER_FILL_COLOR)
HEADER_ALIGNMENT = Alignment(horizontal="center", vertical="center", wrap_text=True)
BODY_ALIGNMENT = Alignment(vertical="top", wrap_text=True)

COLUMN_WIDTHS: tuple[int, ...] = (22, 10, 28, 16, 60, 32, 40)


def write_workbook(rows: list[FinalRow], output_path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_NAME

    ws.append(list(WORKBOOK_COLUMNS))
    for col_idx in range(1, len(WORKBOOK_COLUMNS) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = HEADER_ALIGNMENT

    for col_idx, width in enumerate(COLUMN_WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    ws.freeze_panes = "A2"

    for row_index, row in enumerate(rows, start=2):
        payload = row.to_workbook_dict()
        for col_idx, column in enumerate(WORKBOOK_COLUMNS, start=1):
            cell = ws.cell(row=row_index, column=col_idx, value=payload[column])
            cell.alignment = BODY_ALIGNMENT

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
