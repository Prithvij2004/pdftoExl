from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from app.schemas.workbook_schema import WorkbookRow


HEADERS = [
    "Section",
    "Sequence",
    "Question Type",
    "English Question/Index Text",
    "English Answer Text",
    "Needs Review",
    "Confidence Score",
]


def write_workbook_rows_to_excel(rows: list[WorkbookRow], output_path: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Assessment"

    header_font = Font(bold=True)
    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    body_alignment = Alignment(vertical="top", wrap_text=True)

    for col, name in enumerate(HEADERS, start=1):
        cell = ws.cell(row=1, column=col, value=name)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment

    for row_idx, row in enumerate(rows, start=2):
        values = [
            row.section,
            row.sequence,
            row.question_type,
            row.question_text,
            row.answer_text,
            row.needs_review,
            row.confidence_score,
        ]
        for col, value in enumerate(values, start=1):
            ws.cell(row=row_idx, column=col, value=value).alignment = body_alignment

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = "A1:G1"
    widths = [30, 10, 18, 70, 55, 14, 16]
    for idx, width in enumerate(widths, start=1):
        ws.column_dimensions[chr(64 + idx)].width = width
    ws.row_dimensions[1].height = 24

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(output_path))
    return output_path
