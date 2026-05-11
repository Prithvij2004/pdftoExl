from __future__ import annotations

import json
from pathlib import Path

import pytest

from app.markdown.normalizer import normalize_markdown
from app.chunking.markdown_chunker import chunk_markdown
from app.excel.excel_writer import HEADERS, write_workbook_rows_to_excel
from app.mapping.workbook_mapper import map_document_to_workbook_rows
from app.parsers.base import BasePDFParser
from app.parsers.parser_factory import ParserFactory
from app.pipeline import pdf_to_markdown
from app.schemas.form_schema import DocumentExtraction, FormItem, FormSection, Option
from app.schemas.workbook_schema import WorkbookRow


class FakeParser(BasePDFParser):
    def parse_to_markdown(self, pdf_path: Path) -> str:
        return "<!-- page: 1 -->\n\n# Page 1\n\nA   B\n\n\n\n- item\n"


def test_parser_factory_returns_pdfplumber_parser() -> None:
    parser = ParserFactory.create("pdfplumber")
    assert parser.__class__.__name__ == "PdfPlumberPDFParser"


def test_parser_factory_rejects_unknown_parser() -> None:
    with pytest.raises(ValueError, match="Unsupported parser: xyz"):
        ParserFactory.create("xyz")


def test_normalize_markdown_removes_excessive_blank_lines() -> None:
    normalized = normalize_markdown("A\n\n\n\nB\n• item\nApplicant Name: ____")
    assert "\n\n\n" not in normalized
    assert "- item" in normalized
    assert "Applicant Name: ____" in normalized


def test_convert_pdf_to_markdown_creates_markdown_and_metadata(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    pdf_path = tmp_path / "input.pdf"
    pdf_path.write_bytes(b"%PDF-1.4\n")

    monkeypatch.setattr(pdf_to_markdown.ParserFactory, "create", lambda parser_name: FakeParser())
    result = pdf_to_markdown.convert_pdf_to_markdown(
        pdf_path,
        parser_name="fake",
        output_dir=tmp_path,
    )

    assert result.markdown_path.exists()
    assert result.metadata_path.exists()
    assert result.markdown_path.parent.name == "markdowns"
    assert "<!-- page: 1 -->" in result.markdown_path.read_text(encoding="utf-8")

    metadata = json.loads(result.metadata_path.read_text(encoding="utf-8"))
    assert metadata["status"] == "success"
    assert metadata["parser"] == "fake"


def test_workbook_mapping_generates_dense_sequences_and_inline_branching() -> None:
    doc = DocumentExtraction(
        document_title="Safety Determination Request Form",
        sections=[
            FormSection(
                section_id="sec_001",
                section_title="Current Living Arrangements",
                items=[
                    FormItem(
                        item_id="item_001",
                        item_type="radio_group",
                        question_text="Applicant residence",
                        confidence=0.9,
                        options=[
                            Option(option_id="opt_001", option_text="Lives alone"),
                            Option(
                                option_id="opt_002",
                                option_text="Other",
                                inline_input={
                                    "label": "Specify",
                                    "question_type": "Text Box",
                                    "source_text": "Other - specify ____",
                                },
                            ),
                        ],
                    )
                ],
            )
        ],
    )

    rows = map_document_to_workbook_rows(doc)

    assert [row.sequence for row in rows] == list(range(1, len(rows) + 1))
    assert rows[1].question_type == "Radio Button"
    assert rows[2].branching_logic == 'Display if Q2 = "Other"'


def test_excel_writer_outputs_expected_columns(tmp_path: Path) -> None:
    out = tmp_path / "out.xlsx"
    rows = [
        WorkbookRow(
            section="Section",
            sequence=1,
            question_type="Display",
            question_text="New Section",
            answer_text="Section",
        )
    ]

    write_workbook_rows_to_excel(rows, out)

    from openpyxl import load_workbook

    wb = load_workbook(out)
    ws = wb.active
    assert [ws.cell(row=1, column=i).value for i in range(1, len(HEADERS) + 1)] == HEADERS
    assert ws.cell(row=2, column=2).value == 1
    assert ws.max_column == 7
    wb.close()


def test_mapper_splits_display_label_and_contextualizes_signature_block() -> None:
    doc = DocumentExtraction(
        sections=[
            FormSection(
                section_id="sec_001",
                section_title="Individual Service Plan - Signature Page",
                items=[
                    FormItem(
                        item_id="display_001",
                        item_type="display",
                        display_text="Freedom of Choice: I choose services.",
                        confidence=0.9,
                    ),
                    FormItem(item_id="name_001", item_type="text_box", question_text="Printed Name", confidence=0.9),
                    FormItem(item_id="sig_001", item_type="signature", question_text="Signature", confidence=0.9),
                    FormItem(item_id="date_001", item_type="date", question_text="Date", confidence=0.9),
                ],
            )
        ]
    )

    rows = map_document_to_workbook_rows(doc)

    assert rows[1].question_text == "Freedom of Choice"
    assert rows[1].answer_text == "Freedom of Choice: I choose services."
    assert rows[2].question_text == "Applicant/Member or Authorized Representative Printed Name:"
    assert rows[3].question_text == "Applicant/Member or Authorized Representative Signature:"
    assert rows[4].question_text == "Applicant/Member or Authorized Representative Date Signed:"


def test_semantic_chunker_keeps_form_blocks_together() -> None:
    markdown = """<!-- page: 1 -->

# Page 1

Safety Determination Request Form

Applicant Name: ____ SSN: ____ DOB: ____

Current Living Arrangements:

Applicant residence:

Lives in own home/apt (alone)

Lives in other’s home—specify relationship ____

Other—specify____

If the applicant would not be able to return to this residence, please explain why:
____

Justification for Safety Determination Request:

The applicant has an approved acuity score of at least five (5) but no more than eight (8).

o Provide a detailed description of the safety concern.

Description of documentation attached: ____

The applicant has an individual acuity score of at least 3 for the mobility or transfer measures.

o Describe how often mobility and/or transfer assistance is needed.

Description of documentation attached: ____

<!-- page: 2 -->

Recent (last 365 days) hospital admissions

| Admit Date | Discharge Date | Reason for Admission |
| --- | --- | --- |
|  |  |  |

Submitting Entity Attestation

Applicant, Member or Authorized Representative:

Printed Name

Signature

Date
"""

    chunks = chunk_markdown(markdown, max_chars=12000)
    texts = [chunk.chunk_text for chunk in chunks]

    living = next(text for text in texts if "Current Living Arrangements" in text)
    assert "Lives in other’s home" in living
    assert "Other—specify" in living
    assert "Justification for Safety Determination Request" not in living

    first_criterion = next(text for text in texts if "approved acuity score" in text)
    assert "Provide a detailed description" in first_criterion
    assert "Description of documentation attached" in first_criterion

    table_chunk = next(text for text in texts if "Recent (last 365 days) hospital admissions" in text)
    assert "| Admit Date | Discharge Date | Reason for Admission |" in table_chunk

    signature_chunk = next(text for text in texts if "Applicant, Member or Authorized Representative" in text)
    assert "Printed Name" in signature_chunk
    assert "Signature" in signature_chunk
    assert "Date" in signature_chunk
