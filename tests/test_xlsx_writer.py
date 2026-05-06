from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from app.services.merge_resolve import WORKBOOK_COLUMNS, FinalRow
from app.services.xlsx_writer import (
    COLUMN_WIDTHS,
    HEADER_FILL_COLOR,
    SHEET_NAME,
    write_workbook,
)


def _row(
    *,
    section: str | None,
    sequence: int,
    question_type: str,
    question_text: str,
    branching_logic: str | None = None,
    answer_text: str | None = None,
    question_rule: str | None = None,
) -> FinalRow:
    return FinalRow(
        section=section,
        sequence=sequence,
        question_rule=question_rule,
        question_type=question_type,
        question_text=question_text,
        branching_logic=branching_logic,
        answer_text=answer_text,
        source_ids=["p1_b01"],
        chunk_id="p01_c01",
        page_numbers=[1],
        confidence=0.9,
    )


def test_write_workbook_emits_styled_headers_and_rows(tmp_path: Path) -> None:
    rows = [
        _row(
            section="Intake",
            sequence=1,
            question_type="Display",
            question_text="New Section",
            answer_text="Intake",
        ),
        _row(
            section="Intake",
            sequence=2,
            question_type="Text Box",
            question_text="Applicant Name",
        ),
        _row(
            section="Intake",
            sequence=3,
            question_type="Radio Button",
            question_text="Marital Status",
            answer_text="Single\nMarried",
            branching_logic='Display if Q2 = "Single"',
        ),
    ]

    output_path = tmp_path / "out.xlsx"
    write_workbook(rows, output_path)

    assert output_path.exists()
    wb = load_workbook(output_path)
    assert wb.sheetnames == [SHEET_NAME]
    ws = wb[SHEET_NAME]

    header_values = [ws.cell(row=1, column=i + 1).value for i in range(len(WORKBOOK_COLUMNS))]
    assert tuple(header_values) == WORKBOOK_COLUMNS

    header_cell = ws.cell(row=1, column=1)
    assert header_cell.font.bold is True
    assert header_cell.font.color.rgb.endswith("FFFFFF")
    assert header_cell.fill.fgColor.rgb.endswith(HEADER_FILL_COLOR)

    assert ws.freeze_panes == "A2"

    for col_idx, width in enumerate(COLUMN_WIDTHS, start=1):
        column_letter = ws.cell(row=1, column=col_idx).column_letter
        assert ws.column_dimensions[column_letter].width == width

    def _norm(values: list[object]) -> list[object]:
        return ["" if v is None else v for v in values]

    second_row = _norm([ws.cell(row=2, column=i + 1).value for i in range(len(WORKBOOK_COLUMNS))])
    assert second_row == ["Intake", 1, "", "Display", "New Section", "", "Intake"]

    third_row = _norm([ws.cell(row=3, column=i + 1).value for i in range(len(WORKBOOK_COLUMNS))])
    assert third_row == ["Intake", 2, "", "Text Box", "Applicant Name", "", ""]

    fourth_row = _norm([ws.cell(row=4, column=i + 1).value for i in range(len(WORKBOOK_COLUMNS))])
    assert fourth_row == [
        "Intake",
        3,
        "",
        "Radio Button",
        "Marital Status",
        'Display if Q2 = "Single"',
        "Single\nMarried",
    ]

    body_cell = ws.cell(row=2, column=5)
    assert body_cell.alignment.wrap_text is True
    assert body_cell.alignment.vertical == "top"


def test_write_workbook_omits_internal_fields(tmp_path: Path) -> None:
    rows = [
        _row(
            section=None,
            sequence=1,
            question_type="Text Box",
            question_text="Field",
        ),
    ]
    output_path = tmp_path / "out.xlsx"
    write_workbook(rows, output_path)

    wb = load_workbook(output_path)
    ws = wb[SHEET_NAME]
    assert ws.max_column == len(WORKBOOK_COLUMNS)
    written = [ws.cell(row=2, column=i + 1).value for i in range(ws.max_column)]
    assert "p01_c01" not in written
    assert "p1_b01" not in written
