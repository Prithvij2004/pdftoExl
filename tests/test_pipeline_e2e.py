from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook

from app.config import PipelineConfig
from app.services.chunk_extractor import (
    ChunkExtractionResult,
    ChunkExtractor,
    ExtractedRow,
)
from app.services.chunker import Chunk
from app.services.layout_parser import RawPdfModel
from app.services.page_preparation import PreparedPdf
from app.services.pipeline import (
    CurrentExtractionPipeline,
    ExtractionOutcome,
    PipelineRunResult,
)


def _config(tmp_path: Path) -> PipelineConfig:
    return PipelineConfig(
        bedrock_region=None,
        bedrock_vision_model_id="vision",
        bedrock_text_model_id="text",
        page_render_dpi=72,
        complexity_uncertain_threshold=0.15,
        max_chunk_input_tokens=3500,
        max_concurrent_bedrock_calls=2,
        extraction_cache_dir=tmp_path / "cache",
    )


@dataclass(frozen=True)
class _StubChunker:
    chunks: list[Chunk]

    def build(
        self,
        parser_artifacts,
        vlm_artifacts,
        semantic_pages=None,
    ) -> list[Chunk]:
        return list(self.chunks)


class _StubChunkExtractor(ChunkExtractor):
    def __init__(self, results_by_chunk: dict[str, ChunkExtractionResult]) -> None:
        self.results_by_chunk = results_by_chunk
        self.calls: list[str] = []

    def extract(self, chunk: Chunk) -> ChunkExtractionResult:
        self.calls.append(chunk.chunk_id)
        return self.results_by_chunk[chunk.chunk_id]


def _chunk(chunk_id: str, order_key: tuple[int, int], section: str | None) -> Chunk:
    return Chunk(
        chunk_id=chunk_id,
        pages=[order_key[0]],
        order_key=order_key,
        section_hint=section,
        source_block_ids=[],
        chunk_type_hint="page",
        markdown="# stub\n",
        structure_hints=[],
    )


def _extracted(
    chunk_id: str,
    *,
    question_type: str,
    question_text: str,
    section: str | None = None,
    answer_text: str | None = None,
    branching_logic: str | None = None,
) -> ExtractedRow:
    return ExtractedRow(
        section=section,
        sequence=None,
        question_rule=None,
        question_type=question_type,
        question_text=question_text,
        branching_logic=branching_logic,
        answer_text=answer_text,
        source_ids=[],
        chunk_id=chunk_id,
        page_numbers=[1],
        confidence=0.9,
    )


def test_extract_to_workbook_drives_full_pipeline(tmp_path: Path) -> None:
    config = _config(tmp_path)

    chunks = [
        _chunk("p01_c01", (1, 1), section="Intake"),
        _chunk("p01_c02", (1, 2), section="Intake"),
    ]
    chunk_results = {
        "p01_c01": ChunkExtractionResult(
            chunk_id="p01_c01",
            rows=[
                _extracted(
                    "p01_c01",
                    question_type="Display",
                    question_text="New Section",
                    section="Intake",
                    answer_text="Intake",
                ),
                _extracted(
                    "p01_c01",
                    question_type="Text Box",
                    question_text="Applicant Name",
                    section="Intake",
                ),
            ],
        ),
        "p01_c02": ChunkExtractionResult(
            chunk_id="p01_c02",
            rows=[
                _extracted(
                    "p01_c02",
                    question_type="Date",
                    question_text="Date of Birth",
                    section="Intake",
                ),
            ],
        ),
    }

    pipeline = CurrentExtractionPipeline(
        config,
        chunker=_StubChunker(chunks),
        chunk_extractor=_StubChunkExtractor(chunk_results),
    )

    fake_run = PipelineRunResult(
        pdf_path=Path("synthetic.pdf"),
        prepared_pdf=PreparedPdf(
            pdf_path=Path("synthetic.pdf"),
            pdf_hash="deadbeef",
            renderer_version="r",
            render_dpi=72,
            pages=[],
        ),
        raw_pdf=RawPdfModel(pdf_path=Path("synthetic.pdf"), parser_version="p", pages=[]),
        semantic_pages=[],
        complexity_routes=[],
        parser_artifacts=[],
        vlm_artifacts=[],
    )
    pipeline.run = lambda pdf_path: fake_run  # type: ignore[assignment]

    pdf_path = tmp_path / "synthetic.pdf"
    pdf_path.write_bytes(b"%PDF-stub")
    xlsx_path = tmp_path / "out.xlsx"

    outcome = pipeline.extract_to_workbook(pdf_path, xlsx_path)

    assert isinstance(outcome, ExtractionOutcome)
    assert outcome.xlsx_path == xlsx_path
    assert xlsx_path.exists()
    assert outcome.run_result.chunks == chunks
    assert len(outcome.run_result.chunk_results) == 2
    assert outcome.final_rows, "expected final rows"

    sequences = [row.sequence for row in outcome.final_rows]
    assert sequences == list(range(1, len(outcome.final_rows) + 1))

    wb = load_workbook(xlsx_path)
    ws = wb["Form"]
    headers = [ws.cell(row=1, column=i + 1).value for i in range(7)]
    assert headers == [
        "Section",
        "Sequence",
        "Question Rule",
        "Question Type",
        "Question Text",
        "Branching Logic",
        "Answer Text",
    ]
    body_question_texts = [ws.cell(row=r, column=5).value for r in range(2, ws.max_row + 1)]
    assert "Applicant Name" in body_question_texts
    assert "Date of Birth" in body_question_texts


def test_run_to_workbook_handles_no_chunks(tmp_path: Path) -> None:
    config = _config(tmp_path)
    pipeline = CurrentExtractionPipeline(
        config,
        chunker=_StubChunker([]),
        chunk_extractor=_StubChunkExtractor({}),
    )
    fake_run = PipelineRunResult(
        pdf_path=Path("synthetic.pdf"),
        prepared_pdf=PreparedPdf(
            pdf_path=Path("synthetic.pdf"),
            pdf_hash="deadbeef",
            renderer_version="r",
            render_dpi=72,
            pages=[],
        ),
        raw_pdf=RawPdfModel(pdf_path=Path("synthetic.pdf"), parser_version="p", pages=[]),
        semantic_pages=[],
        complexity_routes=[],
        parser_artifacts=[],
        vlm_artifacts=[],
    )
    pipeline.run = lambda pdf_path: fake_run  # type: ignore[assignment]

    xlsx_path = tmp_path / "empty.xlsx"
    outcome = pipeline.extract_to_workbook(tmp_path / "synthetic.pdf", xlsx_path)
    assert outcome.final_rows == []
    assert xlsx_path.exists()
